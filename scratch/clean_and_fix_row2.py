import openpyxl
from openpyxl.styles import Border, PatternFill
from openpyxl.utils import get_column_letter
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'd:\HA\TAKENOTE CÁC CUỘC THI PROJECT TIN314.xlsx')
ws_rev = wb['Phân tích và chỉnh sửa']
ws_main = wb['Tổng hợp ']

# Xóa các dòng rác 19-35
for r in range(19, 36):
    for c in range(1, 8):
        ws_rev.cell(r, c).value = None
        ws_rev.cell(r, c).border = Border()
        ws_rev.cell(r, c).fill = PatternFill(fill_type=None)

# Đặt lại công thức chính xác cho Hàng 2 của Tổng hợp trỏ về dòng 6-18
cols_mapping = [
    (1, 6),   # Tên cuộc thi -> B6
    (2, 7),   # Lĩnh vực -> B7
    (3, 8),   # Đơn vị tổ chức -> B8
    (4, 9),   # Mô tả ngắn gọn -> B9
    (5, 10),  # Đối tượng tham gia -> B10
    (6, 11),  # Hình thức tham gia -> B11
    (7, 12),  # Cơ cấu vòng thi -> B12
    (8, 13),  # Cơ cấu giải thưởng -> B13
    (9, 14),  # Các mốc thời gian -> B14
    (10, 15), # Mùa thi -> B15
    (11, 16), # Deadline đăng ký -> B16
    (12, 17), # Trạng thái -> B17
    (13, 18)  # Link nguồn chính thức -> B18
]

for col_idx, rev_row in cols_mapping:
    ws_main.cell(2, col_idx).value = f"='Phân tích và chỉnh sửa'!B{rev_row}"
    ws_rev.cell(rev_row, 6).value = f"='Tổng hợp '!{get_column_letter(col_idx)}2"

wb.save(r'd:\HA\TAKENOTE CÁC CUỘC THI PROJECT TIN314.xlsx')
print('Cleaned rows 19-35 and set Row 2 formulas successfully!')
