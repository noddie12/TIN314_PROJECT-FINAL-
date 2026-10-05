import openpyxl
from openpyxl.utils import get_column_letter
import sys

sys.stdout.reconfigure(encoding='utf-8')
file_path = r'd:\HA\TAKENOTE CÁC CUỘC THI PROJECT TIN314.xlsx'
wb = openpyxl.load_workbook(file_path)
ws_main = wb['Tổng hợp ']
ws_rev = wb['Phân tích và chỉnh sửa']

# Thiết lập chế độ tính toán tự động ép buộc khi mở file trong Excel
if hasattr(wb, 'calculation') and wb.calculation:
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

# Liên kết chính xác 13 cột của cuộc thi Khởi nghiệp cùng Kawai (Hàng #2 trên Tổng hợp <-> Dòng 6-18 trên Phân tích và chỉnh sửa)
cols_map = [
    (1, 6),   # Tên cuộc thi -> B6
    (2, 7),   # Lĩnh vực -> B7
    (3, 8),   # Đơn vị tổ chức -> B8
    (4, 9),   # Mô tả ngắn gọn -> B9
    (5, 10),  # Đối tượng tham gia -> B10
    (6, 11),  # Hình thức tham gia -> B11
    (7, 12),  # Cơ cấu vòng thi -> B12
    (8, 13),  # Cơ cấu giải thưởng -> B13
    (9, 14),  # Các mốc thời gian -> B14
    (10, 15), # Mùa thi -> B15
    (11, 16), # Deadline đăng ký -> B16
    (12, 17), # Trạng thái -> B17
    (13, 18)  # Link nguồn chính thức -> B18
]

for col_idx, rev_row in cols_map:
    formula = f"='Phân tích và chỉnh sửa'!B{rev_row}"
    ws_main.cell(2, col_idx).value = formula
    ws_rev.cell(rev_row, 6).value = f"='Tổng hợp '!{get_column_letter(col_idx)}2"

wb.save(file_path)
print("SUCCESS_SAVED")
