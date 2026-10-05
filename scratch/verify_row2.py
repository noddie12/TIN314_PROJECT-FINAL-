import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook(r'd:\HA\TAKENOTE CÁC CUỘC THI PROJECT TIN314.xlsx', data_only=False)
ws_main = wb['Tổng hợp ']
ws_rev = wb['Phân tích và chỉnh sửa']

print('=== KIỂM TRA LIÊN KẾT HÀNG #2 TRANG TỔNG HỢP ===')
all_ok = True
for c in range(1, 14):
    h = str(ws_main.cell(1, c).value or "").strip()
    v = str(ws_main.cell(2, c).value or "")
    rev_r = c + 5  # Col 1 -> B6, Col 2 -> B7, ..., Col 13 -> B18
    expected = f"='Phân tích và chỉnh sửa'!B{rev_r}"
    rev_val = ws_rev.cell(rev_r, 2).value
    col_f_val = ws_rev.cell(rev_r, 6).value
    print(f"Col {c:2d} ({h:<22}): Công thức={v:<32} | Giá trị B{rev_r}={str(rev_val)[:30]}")
    if v != expected:
        all_ok = False

print(f"\nTrạng thái tính toán tự động: calcMode={wb.calculation.calcMode}, fullCalcOnLoad={wb.calculation.fullCalcOnLoad}")
if all_ok:
    print("✅ TẤT CẢ 13 CỘT ĐỀU ĐÃ ĐƯỢC LIÊN KẾT CHUẨN XÁC 100%!")
