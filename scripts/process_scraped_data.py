#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
===================================================================================
COMPETEMATE: CHƯƠNG TRÌNH PHÂN TÍCH DỮ LIỆU CÀO FACEBOOK & GHI EXCEL TỰ ĐỘNG
===================================================================================
Chương trình tối ưu hoá việc tận dụng hệ thống AI đa tầng (Multi-model AI Pool):
- Tự động quét và khớp các bài đăng Facebook đã cào về trong thư mục data/raw_posts/
- Gọi AI Engine (hỗ trợ tự động fallback khi timeout/lỗi hạn mức, Circuit Breaker)
- Trích xuất thông tin chi tiết (Competition DNA & 13 cột nghiệp vụ)
- Bảo toàn tuyệt đối danh mục định sẵn (Column 2) theo quy chuẩn mẫu TAKENOTE
- Ghi dữ liệu vào trang 'Phân tích và chỉnh sửa' và tạo liên kết công thức 2 chiều sang trang 'Tổng hợp '
- Đồng bộ vào cơ sở dữ liệu SQLite và JSON web
===================================================================================
"""

import sys
import os
import re
import json
import time
import argparse
from pathlib import Path
from typing import List, Dict, Any, Optional

# Đảm bảo UTF-8 trên Windows console
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Thêm thư mục gốc vào sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import openpyxl
from pipeline.config import EXCEL_FILE, DATA_DIR, RAW_POSTS_DIR, AI_MODEL, DEFAULT_FALLBACK_MODELS
from pipeline.ai_extractor import (
    AICompetitionExtractor,
    is_model_quarantined,
    quarantine_model,
    reset_model_quarantine,
    get_quarantined_models_summary
)
from pipeline.excel_sync import ExcelSyncManager, REVIEW_SHEET_NAME
from pipeline.db_syncer import DatabaseSyncer
from pipeline.models import CompetitionDNA

# ==============================================================================
# HÀM HỖ TRỢ CHUẨN HOÁ VÀ TÌM KIẾM
# ==============================================================================
def clean_str(s: Any) -> str:
    """Loại bỏ ký tự đặc biệt, đưa về chữ thường để so khớp chính xác."""
    if not s:
        return ""
    return re.sub(r"[^a-zA-Z0-9]", "", str(s).lower())

def extract_slug_from_url(url: str) -> str:
    """Tách slug nhận diện từ URL Facebook hoặc Website."""
    if not url:
        return ""
    clean = re.sub(r"\?.*$", "", str(url)).rstrip("/").split("/")[-1]
    return clean

def load_competitions_metadata() -> List[Dict[str, Any]]:
    """Nạp danh sách chuẩn 102 cuộc thi và danh mục định sẵn."""
    links_file = DATA_DIR / "competitions_links.json"
    cat_file = DATA_DIR / "predefined_categories.json"
    
    comp_links = []
    if links_file.exists():
        with open(links_file, "r", encoding="utf-8") as f:
            comp_links = json.load(f)
            
    cat_by_row = {}
    if cat_file.exists():
        with open(cat_file, "r", encoding="utf-8") as f:
            cat_by_row = json.load(f).get("by_row", {})
            
    for item in comp_links:
        stt = item.get("stt", 0)
        target_row = stt + 1
        item["target_row"] = target_row
        if str(target_row) in cat_by_row:
            item["category"] = cat_by_row[str(target_row)]
        elif target_row in cat_by_row:
            item["category"] = cat_by_row[target_row]
            
    return comp_links

# ==============================================================================
# TRÌNH QUẢN LÝ QUÉT & PHÂN TÍCH (SCANNER & RUNNER)
# ==============================================================================
class ScrapedCompetitionManager:
    def __init__(self, excel_path: Optional[Path] = None):
        self.excel_file = excel_path or EXCEL_FILE
        self.excel_mgr = ExcelSyncManager(self.excel_file)
        self.db_syncer = DatabaseSyncer()
        self.metadata = load_competitions_metadata()
        self.extractor = AICompetitionExtractor(model_name="kc/kilo-auto/free")
        
    def scan_all(self) -> List[Dict[str, Any]]:
        """
        Quét toàn bộ thư mục data/raw_posts/, đối chiếu với danh sách cuộc thi
        và kiểm tra trạng thái hiện tại trong file Excel.
        """
        raw_files = list(RAW_POSTS_DIR.glob("*.json"))
        
        # Mở workbook để kiểm tra tình trạng từng hàng
        wb = openpyxl.load_workbook(self.excel_file, data_only=False)
        ws = wb["Tổng hợp "] if "Tổng hợp " in wb.sheetnames else wb.active
        
        results = []
        for comp in self.metadata:
            stt = comp.get("stt", 0)
            row = comp.get("target_row", stt + 1)
            name = comp.get("name", "")
            cat = comp.get("category", "")
            url = comp.get("url", "")
            
            slug = extract_slug_from_url(url)
            norm_slug = clean_str(slug)
            norm_name = clean_str(name)
            
            # Tìm file raw posts khớp
            matched_files = []
            for rf in raw_files:
                if rf.stat().st_size < 500:
                    continue  # Bỏ qua file rỗng hoặc lỗi cào
                rf_clean = clean_str(rf.stem)
                if norm_slug and (norm_slug in rf_clean or rf_clean.startswith(norm_slug)):
                    matched_files.append(rf)
                elif len(norm_name) >= 6 and (norm_name in rf_clean or rf_clean in norm_name):
                    matched_files.append(rf)
                    
            latest_file = max(matched_files, key=os.path.getmtime) if matched_files else None
            
            # Kiểm tra trạng thái dữ liệu trong Excel
            # Dòng coi là "Đã có dữ liệu" nếu Cột 3 (Đơn vị) hoặc Cột 4 (Mô tả) có nội dung
            org_cell = ws.cell(row, 3)
            desc_cell = ws.cell(row, 4)
            org_val = self.excel_mgr._resolve_cell_value(org_cell, wb)
            desc_val = self.excel_mgr._resolve_cell_value(desc_cell, wb)
            
            has_full_excel = bool((org_val and str(org_val).strip()) or (desc_val and str(desc_val).strip()))
            is_formula = str(ws.cell(row, 1).value or "").startswith(f"='{REVIEW_SHEET_NAME}'")
            
            post_count = 0
            file_size_kb = 0
            if latest_file:
                file_size_kb = round(latest_file.stat().st_size / 1024, 1)
                try:
                    with open(latest_file, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        post_count = len(data) if isinstance(data, list) else 1
                except Exception:
                    pass
                    
            results.append({
                "stt": stt,
                "target_row": row,
                "name": name,
                "category": cat,
                "url": url,
                "slug": slug,
                "has_scraped_file": bool(latest_file),
                "scraped_file": latest_file,
                "post_count": post_count,
                "file_size_kb": file_size_kb,
                "has_excel_data": has_full_excel,
                "is_formula_linked": is_formula
            })
            
        return results

    def print_status_table(self, items: Optional[List[Dict[str, Any]]] = None):
        """In bảng trạng thái tổng thể các cuộc thi."""
        if items is None:
            items = self.scan_all()
            
        print("\n" + "=" * 90)
        print("📊 BẢNG TỔNG HỢP TRẠNG THÁI CÀO DỮ LIỆU & GHI FILE EXCEL")
        print("=" * 90)
        print(f"{'Hàng':<6} | {'Tên cuộc thi':<32} | {'Lĩnh vực':<16} | {'Dữ liệu cào':<16} | {'Excel Status'}")
        print("-" * 90)
        
        scraped_count = 0
        excel_count = 0
        pending_count = 0
        
        for it in items:
            if it["has_scraped_file"]:
                scraped_count += 1
                scrap_str = f"✅ {it['post_count']} bài ({it['file_size_kb']}KB)"
            else:
                scrap_str = "⚪ Chưa có file"
                
            if it["has_excel_data"]:
                excel_count += 1
                excel_str = "🟢 Đã có dữ liệu (Liên kết công thức)" if it["is_formula_linked"] else "🟢 Đã có dữ liệu (Tĩnh)"
            elif it["has_scraped_file"]:
                pending_count += 1
                excel_str = "⏳ ĐANG CHỜ PHÂN TÍCH"
            else:
                excel_str = "⚪ Hàng mẫu trống"
                
            # Chỉ in các cuộc thi có file cào hoặc đã có trong Excel để bảng gọn gàng
            if it["has_scraped_file"] or it["has_excel_data"]:
                print(f"#{it['target_row']:<5} | {it['name'][:30]:<32} | {it['category'][:14]:<16} | {scrap_str:<16} | {excel_str}")
                
        print("-" * 90)
        print(f"📌 TỔNG KẾT:")
        print(f"   - Tổng số cuộc thi đã cào thành công: {scraped_count}")
        print(f"   - Cuộc thi đã được phân tích và ghi Excel: {excel_count}")
        print(f"   - Cuộc thi ĐÃ CÀO nhưng ĐANG CHỜ PHÂN TÍCH VÀO EXCEL: {pending_count}")
        print("=" * 90 + "\n")

    def process_single(
        self,
        item: Dict[str, Any],
        force: bool = False,
        verbose: bool = True
    ) -> bool:
        """
        Thực hiện phân tích bài đăng đã cào cho 1 cuộc thi và ghi vào Excel.
        """
        row = item["target_row"]
        name = item["name"]
        cat = item["category"]
        raw_file = item["scraped_file"]
        url = item["url"]
        
        if not raw_file or not Path(raw_file).exists():
            if verbose:
                print(f"❌ [Hàng #{row} - {name}] Không tìm thấy file cào trong data/raw_posts/!")
            return False
            
        if item["has_excel_data"] and not force:
            if verbose:
                print(f"⏭️ [Hàng #{row} - {name}] Đã có dữ liệu trong Excel. (Dùng cờ --force để ghi đè)")
            return True
            
        print("\n" + "=" * 70)
        print(f"🚀 BẮT ĐẦU XỬ LÝ: #{row} - {name}")
        print(f"📁 File nguồn: {raw_file.name} ({item['post_count']} bài đăng, {item['file_size_kb']} KB)")
        print(f"🏷️ Lĩnh vực định sẵn: {cat}")
        print("=" * 70)
        
        # 1. Đọc nội dung file JSON
        try:
            with open(raw_file, "r", encoding="utf-8") as f:
                posts = json.load(f)
            if not isinstance(posts, list) or len(posts) == 0:
                print(f"⚠️ File cào {raw_file.name} không có bài viết hợp lệ!")
                return False
        except Exception as e:
            print(f"❌ Lỗi đọc file JSON: {e}")
            return False
            
        # 2. Bước 1: Trích xuất Competition DNA qua AI Engine
        print(f"\n[AI 1/2] Đang trích xuất Competition DNA (Name, Season, Organizer, Deadlines, Prizes)...")
        t0 = time.time()
        try:
            dna = self.extractor.extract_competition_dna(
                posts=posts,
                fallback_name=name,
                fallback_organizer=f"BTC {name}",
                fallback_url=url
            )
            elapsed_dna = time.time() - t0
            print(f"✅ DNA thành công ({elapsed_dna:.1f}s):")
            print(f"   - Tên cuộc thi : {dna.name}")
            print(f"   - Mùa thi      : {dna.season or 'Không rõ'}")
            print(f"   - Đơn vị       : {dna.organizer}")
            print(f"   - Hạn đăng ký  : {dna.deadline or 'Chưa xác định'}")
            print(f"   - Trạng thái   : {dna.status}")
        except Exception as e:
            print(f"❌ Không thể trích xuất DNA bằng AI: {e}")
            return False
            
        # 3. Bước 2: Ánh xạ 13 cột nghiệp vụ cho Excel
        print(f"\n[AI 2/2] Đang ánh xạ dữ liệu vào 13 cột Excel...")
        cols = self.excel_mgr.get_schema()["columns"]
        mapped_data = None
        t1 = time.time()
        try:
            mapped_data = self.extractor.extract_custom_mapping(
                posts=posts,
                target_columns=cols,
                fallback_name=dna.name,
                fallback_organizer=dna.organizer,
                fallback_url=url
            )
            elapsed_map = time.time() - t1
            print(f"✅ Ánh xạ thành công {len(mapped_data)} cột ({elapsed_map:.1f}s)!")
        except Exception as e:
            print(f"⚠️ Trích xuất custom mapping gặp trục trặc: {e}")
            print(f"🔄 Tự động chuyển đổi từ Competition DNA để điền đầy đủ dữ liệu vào Excel...")
            mapped_data = dna.to_excel_dict()
            
        # Đảm bảo giữ nguyên Lĩnh vực định sẵn trong mapped_data nếu có
        if cat:
            mapped_data["Lĩnh vực "] = cat
            mapped_data["Lĩnh vực"] = cat
            
        # 4. Bước 3: Ghi vào Excel tại hàng được chỉ định
        print(f"\n[Excel Sync] Đang ghi vào hàng #{row} của {self.excel_file.name}...")
        try:
            saved_row = self.excel_mgr.update_competition(
                dna=dna,
                custom_data=mapped_data,
                target_row=row
            )
            print(f"✅ Đã ghi thành công vào Excel tại hàng #{saved_row}!")
            print(f"   - Giữ nguyên Lĩnh vực: '{cat}'")
            print(f"   - Tạo khối dữ liệu tại trang '{REVIEW_SHEET_NAME}'")
            print(f"   - Tạo công thức liên kết 2 chiều trên trang 'Tổng hợp '")
        except Exception as e:
            print(f"❌ Lỗi ghi file Excel: {e}")
            return False
            
        # 5. Bước 4: Đồng bộ vào Database & JSON Web
        print(f"[DB Sync] Đang đồng bộ vào SQLite và Web JSON...")
        try:
            self.db_syncer.sync_competition(dna)
            print(f"✅ Đã đồng bộ Database thành công!")
        except Exception as e:
            print(f"⚠️ Lỗi đồng bộ DB (không ảnh hưởng file Excel): {e}")
            
        print(f"\n🎉 HOÀN TẤT XỬ LÝ #{row} - {dna.name}!")
        return True

    def process_all_pending(self, force: bool = False):
        """Tự động phân tích tất cả các cuộc thi đã cào nhưng chưa có trong Excel."""
        items = self.scan_all()
        pending = [it for it in items if it["has_scraped_file"] and (not it["has_excel_data"] or force)]
        
        if not pending:
            print("\n✨ Tuyệt vời! Tất cả các cuộc thi có file cào đều đã được phân tích và ghi vào Excel.")
            return
            
        print("\n" + "=" * 70)
        print(f"📋 PHÁT HIỆN {len(pending)} CUỘC THI CẦN PHÂN TÍCH VÀ GHI EXCEL:")
        print("=" * 70)
        for idx, it in enumerate(pending, 1):
            print(f"[{idx:02d}] Hàng #{it['target_row']:<3} | {it['name']:<35} | {it['file_size_kb']} KB ({it['post_count']} bài)")
        print("=" * 70)
        
        success_count = 0
        fail_count = 0
        
        for idx, it in enumerate(pending, 1):
            print(f"\n>>> [{idx}/{len(pending)}] TIẾN TRÌNH XỬ LÝ #{it['target_row']} - {it['name']} <<<")
            ok = self.process_single(it, force=force)
            if ok:
                success_count += 1
            else:
                fail_count += 1
            # Nghỉ ngắn giữa các request để bảo đảm an toàn kết nối
            time.sleep(2)
            
        print("\n" + "=" * 70)
        print("🏁 HOÀN TẤT QUÁ TRÌNH TỰ ĐỘNG PHÂN TÍCH TOÀN BỘ!")
        print(f"   - Thành công : {success_count} cuộc thi")
        print(f"   - Thất bại   : {fail_count} cuộc thi")
        print("=" * 70 + "\n")

    def process_by_row(self, row_numbers: List[int], force: bool = False):
        """Xử lý các cuộc thi theo số hàng Excel cụ thể."""
        items = self.scan_all()
        items_by_row = {it["target_row"]: it for it in items}
        
        for r in row_numbers:
            if r not in items_by_row:
                print(f"❌ Không tìm thấy cuộc thi nào tại hàng #{r}!")
                continue
            it = items_by_row[r]
            if not it["has_scraped_file"]:
                print(f"❌ Cuộc thi #{r} ({it['name']}) chưa có file cào bài viết trong data/raw_posts/!")
                continue
            self.process_single(it, force=force)

    def process_by_url_or_slug(self, query: str, force: bool = False):
        """Xử lý cuộc thi theo link URL hoặc từ khóa tên / slug."""
        items = self.scan_all()
        clean_q = clean_str(query)
        matched = []
        for it in items:
            if clean_q in clean_str(it["slug"]) or clean_q in clean_str(it["name"]) or clean_q in clean_str(it["url"]):
                matched.append(it)
                
        if not matched:
            print(f"❌ Không tìm thấy cuộc thi nào khớp với '{query}'!")
            return
            
        for it in matched:
            if not it["has_scraped_file"]:
                print(f"❌ Cuộc thi #{it['target_row']} ({it['name']}) chưa có file cào trong data/raw_posts/!")
                continue
            self.process_single(it, force=force)

    def process_custom_file(self, file_path_str: str, target_row: Optional[int] = None, force: bool = True):
        """Phân tích một file JSON bất kỳ."""
        p = Path(file_path_str)
        if not p.is_absolute():
            p = BASE_DIR / p
            if not p.exists():
                p = RAW_POSTS_DIR / file_path_str
                
        if not p.exists():
            print(f"❌ Không tìm thấy file: {file_path_str}")
            return
            
        # Thử tìm cuộc thi khớp với file
        items = self.scan_all()
        matched_item = None
        for it in items:
            if it["scraped_file"] and it["scraped_file"].name == p.name:
                matched_item = it
                break
                
        if matched_item:
            if target_row:
                matched_item["target_row"] = target_row
            self.process_single(matched_item, force=force)
        else:
            # Tạo item tạm
            target_r = target_row or 2
            item = {
                "target_row": target_r,
                "name": p.stem.split("_")[0],
                "category": "Tài chính",
                "scraped_file": p,
                "url": "",
                "post_count": 0,
                "file_size_kb": round(p.stat().st_size / 1024, 1),
                "has_excel_data": False
            }
            self.process_single(item, force=force)

# ==============================================================================
# MENU TƯƠNG TÁC NGƯỜI DÙNG (INTERACTIVE MENU)
# ==============================================================================
def run_interactive_menu():
    manager = ScrapedCompetitionManager()
    
    while True:
        items = manager.scan_all()
        scraped_items = [it for it in items if it["has_scraped_file"]]
        pending_items = [it for it in scraped_items if not it["has_excel_data"]]
        
        print("\n" + "=" * 72)
        print("🏆 COMPETEMATE: PHÂN TÍCH DỮ LIỆU CÀO FACEBOOK & GHI EXCEL TỰ ĐỘNG")
        print("=" * 72)
        print(f"📊 TÌNH TRẠNG HIỆN TẠI:")
        print(f"   • Số cuộc thi đã cào dữ liệu thô (JSON): {len(scraped_items)}")
        print(f"   • Đã có đầy đủ dữ liệu trong Excel     : {len(scraped_items) - len(pending_items)}")
        print(f"   • ĐANG CHỜ PHÂN TÍCH VÀ GHI VÀO EXCEL  : \033[93m{len(pending_items)}\033[0m cuộc thi")
        print("-" * 72)
        print(" [1] ⚡ Tự động phân tích TẤT CẢ các cuộc thi đang chờ (Khuyên dùng)")
        print(" [2] 🎯 Chọn 1 cuộc thi cụ thể từ danh sách các cuộc thi đã cào")
        print(" [3] 🔢 Nhập STT hàng Excel hoặc Tên / Link Facebook để xử lý")
        print(" [4] 📁 Phân tích trực tiếp 1 file JSON trong thư mục data/raw_posts")
        print(" [5] 📋 Xem bảng trạng thái tổng hợp (Cào vs Excel)")
        print(" [6] 🔄 Xóa cách ly Circuit Breaker (Reset toàn bộ model AI)")
        print(" [0] 🚪 Thoát")
        print("=" * 72)
        
        try:
            choice = input("👉 Lựa chọn của bạn [0-6]: ").strip()
        except (KeyboardInterrupt, EOFError):
            print("\nĐã thoát chương trình.")
            break
            
        if choice == "1":
            manager.process_all_pending(force=False)
            
        elif choice == "2":
            if not scraped_items:
                print("Chưa có file cào nào trong thư mục data/raw_posts/!")
                continue
            print("\n📋 DANH SÁCH CÁC CUỘC THI ĐÃ CÀO BÀI ĐĂNG:")
            print("-" * 72)
            for idx, it in enumerate(scraped_items, 1):
                status_tag = "⏳ [CHỜ PHÂN TÍCH]" if not it["has_excel_data"] else "🟢 [ĐÃ CÓ EXCEL]"
                print(f" [{idx:02d}] Hàng #{it['target_row']:<3} | {it['name']:<35} | {status_tag}")
            print("-" * 72)
            sel = input("👉 Nhập số thứ tự [1-{}]: ".format(len(scraped_items))).strip()
            if sel.isdigit() and 1 <= int(sel) <= len(scraped_items):
                selected_item = scraped_items[int(sel) - 1]
                force_opt = False
                if selected_item["has_excel_data"]:
                    f_in = input("Cuộc thi này đã có trong Excel. Bạn có muốn ghi đè? (y/N): ").strip().lower()
                    force_opt = f_in in ["y", "yes"]
                manager.process_single(selected_item, force=force_opt)
            else:
                print("Lựa chọn không hợp lệ!")
                
        elif choice == "3":
            query = input("👉 Nhập số hàng (VD: 62) hoặc Tên / Link Facebook: ").strip()
            if not query:
                continue
            if query.isdigit():
                manager.process_by_row([int(query)], force=True)
            else:
                manager.process_by_url_or_slug(query, force=True)
                
        elif choice == "4":
            fname = input("👉 Nhập tên file JSON (trong data/raw_posts/): ").strip()
            if not fname:
                continue
            row_in = input("👉 Nhập số hàng Excel đích (để trống nếu muốn tự động khớp): ").strip()
            r_num = int(row_in) if row_in.isdigit() else None
            manager.process_custom_file(fname, target_row=r_num, force=True)
            
        elif choice == "5":
            manager.print_status_table(items)
            
        elif choice == "6":
            reset_model_quarantine()
            print("✅ Đã xóa toàn bộ khóa cách ly Circuit Breaker cho tất cả model AI!")
            
        elif choice == "0":
            print("Tạm biệt!")
            break
        else:
            print("Lựa chọn không hợp lệ, vui lòng chọn lại.")

# ==============================================================================
# HÀM CHÍNH & XỬ LÝ THAM SỐ DÒNG LỆNH (CLI ARGUMENTS)
# ==============================================================================
def main():
    parser = argparse.ArgumentParser(
        description="CompeteMate: Phân tích bài đăng Facebook đã cào bằng AI và ghi vào Excel"
    )
    parser.add_argument("--all", "-a", action="store_true", help="Tự động phân tích tất cả các cuộc thi đang chờ")
    parser.add_argument("--row", "-r", type=int, nargs="+", help="Xử lý theo số hàng Excel (Ví dụ: -r 62 65 67)")
    parser.add_argument("--url", "-u", type=str, help="Xử lý theo Link Facebook hoặc Slug cuộc thi")
    parser.add_argument("--file", "-f", type=str, help="Phân tích trực tiếp 1 file JSON bài đăng thô")
    parser.add_argument("--status", "-s", action="store_true", help="Hiển thị bảng trạng thái tổng quát")
    parser.add_argument("--force", action="store_true", help="Ghi đè nếu hàng đã có dữ liệu")
    parser.add_argument("--reset-cb", action="store_true", help="Xóa khóa cách ly Circuit Breaker")
    
    args = parser.parse_args()
    manager = ScrapedCompetitionManager()
    
    if args.reset_cb:
        reset_model_quarantine()
        print("✅ Đã reset Circuit Breaker.")
        if len(sys.argv) == 2:
            return
            
    if args.status:
        manager.print_status_table()
        return
        
    if args.all:
        manager.process_all_pending(force=args.force)
        return
        
    if args.row:
        manager.process_by_row(args.row, force=args.force)
        return
        
    if args.url:
        manager.process_by_url_or_slug(args.url, force=args.force)
        return
        
    if args.file:
        manager.process_custom_file(args.file, force=args.force)
        return
        
    # Không có tham số -> Chạy giao diện Menu tương tác
    run_interactive_menu()

if __name__ == "__main__":
    main()
